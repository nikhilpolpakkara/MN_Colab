from openpyxl import load_workbook
from DBOps import ExcelOps, MongoDBOps
from DataTools import StringTools, DfTools
from pymongo import MongoClient
from DBOps import TxtOps

class ExcelDataset:
    def __init__(self, excel_file_path):
        self.excel_file_path = excel_file_path
        self.dataset_name = StringTools.get_file_name_without_extension(self.excel_file_path)
        self.wb = load_workbook(excel_file_path)
        self.svp_ws = self.wb["SVP"]
        self.maps_ws = self.wb["MAPS"]
        self.curves_ws = self.wb["CURVES"]
        self.svp_df = None
        self.client = None
        self.datahandler = None

    def load_svp_df(self):
        ref = self.svp_ws.tables["SVP"].ref
        self.svp_df = ExcelOps.get_df_from_cell_reference(sheet=self.svp_ws, ref=ref)

    def get_svp_documents(self):
        df = self.svp_df
        df = df.rename(columns={"NAME": "name", "VALUE": "value"})
        df["function"] = ""
        df["sub_function"] = ""
        df["category"] = ""
        df["variable_type"] = "SVP"
        df['value'] = df['value'].apply(DfTools.convert_to_float_or_str)
        svp_documents = df.to_dict(orient="records")
        return svp_documents

    def get_map_document(self, map_name):
        map_table = self.maps_ws.tables[map_name]
        map_df = ExcelOps.get_df_from_cell_reference(self.maps_ws, map_table.ref)
        map_df = map_df.dropna()
        try:
            map_df.columns = map_df.columns.astype("float")
        except Exception as e:
            print(e)
            pass
        map_df.reset_index(inplace=True)
        map_df = map_df.melt(id_vars='index', var_name='y', value_name='value')
        map_df.columns = ["x", "y", "value"]
        doc = {
            "name": map_name,
            "value": map_df.to_dict(orient="records"),
            "function": "",
            "sub_function": "",
            "category": "",
            "variable_type": "MAP"
        }

        return doc

    def get_curve_document(self, curve_name):
        curve_table = self.curves_ws.tables[curve_name]
        curve_df = ExcelOps.get_df_from_cell_reference(self.curves_ws, curve_table.ref)
        curve_df = curve_df.dropna()
        try:
            curve_df.columns = curve_df.columns.astype("float")
        except Exception as e:
            print(e)
            pass
        curve_df = curve_df.melt(var_name='x', value_name='value')
        doc = {
            "name": curve_name,
            "value": curve_df.to_dict(orient="records"),
            "function": "",
            "sub_function": "",
            "category": "",
            "variable_type": "CURVE"
        }
        return doc

    def load_mongo_client(self, client):
        self.client = client
        self.datahandler = MongoDBOps.MongoDBHandler(client)
        self.datahandler.load_database("CAL")
        self.datahandler.load_collection("DATASETS")

    def check_dataset_document(self):
        dataset_list = self.datahandler.get_field_values_from_level_1_collection(
            field_names=["dataset_id"]
        )
        if self.dataset_name in dataset_list:
            # TODO: if the dataset already exists in the database
            return True
        else:
            doc = {
                "dataset_id": self.dataset_name,
                "hex_data": []
            }
            self.datahandler.add_document_to_collection(doc)
            return False

    def transfer_svp_to_mongodb(self):
        self.load_svp_df()
        self.datahandler.append_list_values_to_level_1_collection_list_field(
            document_filter={"dataset_id": self.dataset_name},
            field_to_update="hex_data",
            new_values=self.get_svp_documents()
        )

    def transfer_maps_to_mongodb(self):
        map_documents = []
        for map_name in self.maps_ws.tables:
            map_doc = self.get_map_document(map_name)
            map_documents.append(map_doc)

        self.datahandler.append_list_values_to_level_1_collection_list_field(
            document_filter={"dataset_id": self.dataset_name},
            field_to_update="hex_data",
            new_values=map_documents
        )

    def transfer_curves_to_mongodb(self):
        curve_documents = []
        for curve_name in self.curves_ws.tables:
            curve_doc = self.get_curve_document(curve_name)
            curve_documents.append(curve_doc)

        self.datahandler.append_list_values_to_level_1_collection_list_field(
            document_filter={"dataset_id": self.dataset_name},
            field_to_update="hex_data",
            new_values=curve_documents
        )


class MongoDBDataset:
    def __init__(self, dataset_id, client):
        self.dataset_id = dataset_id
        self.client = client
        self.handler = MongoDBOps.MongoDBHandler(client)
        self.handler.load_database("CAL")
        self.handler.load_collection("DATASETS")

    def get_variable(self, variable_name):
        dataset_filter = {"dataset_id": self.dataset_id}
        hex_data_field = "hex_data"
        variable_filter = {"field_name": "name", "field_value": variable_name}
        variable_doc = self.handler.get_document_from_level_2_collection(
            nested_array=hex_data_field,
            collection_filter=dataset_filter,
            array_document_filter=variable_filter
        )
        return variable_doc

class DCMdataset:
    def __init__(self, dcm_file_path, ems = "BOSCH"):
        self.dcm_path = dcm_file_path
        self.svp = []
        self.curve = []
        self.map = []
        self.ems = ems
        self.characteristic_identifier : dict

    def detect_ems(self):
        pass
    def get_characteristic_dcm_variables(self):
        """
        Variable names to identify and sort DCM
        :return:
        """
        if self.ems == "BOSCH":
            self.characteristic_identifier = {"map":"KENNFELD",
                                              "curve":"KENNLINIE",
                                              "grouped_map":"GRUPPENKENNFELD",
                                              "grouped_curve":"GRUPPENKENNLINIE",
                                              "svp":"FESTWERT"}

    def extract_curve(self, lines):
        stx_values = []
        wert_values = []

        for line in lines:
            if line.startswith("ST/X"):
                # Extract ST/X values and remove the word "ST/X"
                stx_line = line.replace("ST/X", "").strip()
                stx_values.extend(map(float, stx_line.split()))
            elif line.startswith("WERT"):
                # Extract WERT values and remove the word "WERT"
                wert_line = line.replace("WERT", "").strip()
                wert_values.extend(map(float, wert_line.split()))

        result_list = [{'x': x, 'value': value} for x, value in zip(stx_values, wert_values)]

        return result_list

    def extract_map(self, lines):
        current_map = []
        current_x_values = None
        current_y_values = None

        for line in lines:
            if line.startswith("ST/X"):
                current_x_values = list(map(float, line.replace("ST/X", "").strip().split()))
            elif line.startswith("ST/Y"):
                current_y_values = float(line.replace("ST/Y", "").strip())
            elif line.startswith("WERT"):
                wert_values = list(map(float, line.replace("WERT", "").strip().split()))

                if current_x_values and current_y_values is not None:
                    for x, value in zip(current_x_values, wert_values):
                        current_map.append({'x': x, 'y': current_y_values, 'value': value})

        return current_map
    def get_content(self):
        dcm_content = TxtOps.extract_content_from_txt(self.dcm_path)
        return dcm_content

    def get_functions(self):
        dcm_content = self.get_content()
        function_names = TxtOps.find_lines_between_keywords(dcm_content,"FUNKTIONEN", "END")
        return function_names

    def read_content(self):
        dcm_content = TxtOps.split_lines_by_end(self.get_content(),"END")[1:]
        for content in dcm_content:
            variable_data = {}
            variable_data['name'] = content[1].split()[1]
            variable_data['function_name'] = content[3].split()[1]
            if content[1].startswith(self.characteristic_identifier['svp']):
                variable_data['value'] = content[5].split()[1]
                self.svp.append(variable_data)
            elif content[1].startswith(self.characteristic_identifier['curve']) or \
                    content[1].startswith(self.characteristic_identifier['grouped_curve']):
                variable_data['data'] = self.extract_curve(content)
                self.curve.append(variable_data)
            elif content[1].startswith(self.characteristic_identifier['map']) or \
                    content[1].startswith(self.characteristic_identifier['grouped_map']):
                variable_data['data'] = self.extract_map(content)
                self.map.append(variable_data)




def transfer_dataset_excel_to_mongodb():
    client = MongoClient("mongodb://localhost:27017")
    dataset_excel_file_path = "../data/T400_5N_dataset_step120.xlsx"
    dataset = ExcelDataset(dataset_excel_file_path)
    dataset.load_mongo_client(client=client)
    if not dataset.check_dataset_document():
        dataset.transfer_svp_to_mongodb()
        dataset.transfer_maps_to_mongodb()
        dataset.transfer_curves_to_mongodb()


if __name__ == "__main__":
    # client = MongoClient("mongodb://localhost:27017")
    #
    # # transfer_dataset_excel_to_mongodb()
    #
    # dataset = MongoDBDataset(dataset_id="T400_5N_dataset_step120", client=client)
    # dataset.get_variable("KFMSWDKQ")
    file_path = r"D:\BAL Projects\01_Misc\MN_Colab\Streamlit\data\DCM_STEP2A_V2___FOR_3W.DCM"
    # with open(file_path, 'r') as file:
    #     file_lines = file.readlines()

    dcm_handler = DCMdataset(file_path)
    dcm_handler.get_characteristic_dcm_variables()
    dcm_handler.read_content()





